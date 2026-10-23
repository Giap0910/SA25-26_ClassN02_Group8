from pathlib import Path
from copy import deepcopy
import json
import re
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

ROOT = Path('C:/Project/Badminton_Stock_Security')
WORK = Path(__file__).parent


def replace_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)


def remove_between(doc, start, end):
    children = list(doc.element.body)
    begin = None
    for i, e in enumerate(children):
        text = ''.join(n.text or '' for n in e.iter(qn('w:t')))
        if e.tag.endswith('}p') and text.startswith(start):
            begin = i
            break
    assert begin is not None, start
    finish = next(i for i in range(begin + 1, len(children))
                  if children[i].tag.endswith('}p')
                  and ''.join(n.text or '' for n in children[i].iter(qn('w:t'))).startswith(end))
    for element in children[begin:finish]:
        doc.element.body.remove(element)


def update_report():
    doc = Document(WORK / 'Baocao.before.docx')
    original_group = doc.tables[0]._tbl.xml
    remove_between(doc, 'Hình 28.', '6.12.')
    remove_between(doc, '6.12.', '6.13 ')
    remove_between(doc, '6.4.', '6.10.')
    changes = {
        'Guard trả 0': 'Guard trả 0 chỉ biểu diễn qua điều kiện; việc đánh giá đúng trạng thái tồn kho dựa vào runner và oracle.',
        'Ranh giới cần phân biệt:': 'stock_json trả mã kết quả cho GUI; CLI của adapter xuất cả trạng thái kho để runner đối chiếu. Coverage của một hàm không đại diện toàn bộ parser hoặc website.',
        'Đã có đặc tả, truy vết': 'Phạm vi báo cáo là Ch4 Coverage và Fuzzing: đặc tả đầu vào, đích C, tạo input, coverage, oracle, sanitizer và hồi quy bản vá. Không trình bày kiểm chứng mô hình hoặc phân tích tĩnh như kết quả của Ch4. Không kiểm thử trực tiếp website, transaction, đồng thời, cơ sở dữ liệu, giá, voucher hoặc PayOS.',
        'Guard chỉ kiểm tra điều kiện trả về': 'stock_reserve_guard có thể làm đích fuzz thuần số, nhưng mã 0 chỉ là qua điều kiện; stock_json và stock_json_case mới đi qua parser và nghiệp vụ. Báo cáo không dùng kết quả sinh test ở tab khác làm phép chạy Ch4.',
        'Bước 5. Concolic': 'Bước 5. Lưu test và kiểm tra hồi quy',
        'Nạp guard vào SecLabFramework, nhận': 'Lưu input, seed, ngân sách và snapshot của từng chiến dịch. Xuất test khi công cụ hỗ trợ, giữ đúng nhãn characterization và đối chiếu kỳ vọng với oracle. Sau sửa mã, chạy ca hồi quy cố định và giữ riêng kết quả trước/sau.',
        'GCC/gcov portable cho fuzz': 'GCC/gcov portable dùng chạy C và đo coverage; LLVM-MinGW dùng ASan+UBSan; Python điều phối và đối chiếu oracle; SecLabFramework trình diễn chiến dịch Ch4. Lượt JSON dài có tối đa 500 input và 120 giây cho mỗi chế độ/seed. Coverage-guided không tự đồng nghĩa với concolic.',
        'LeakSanitizer không chạy': 'LeakSanitizer không được bật trong các lượt sanitizer đã báo cáo. ASan+UBSan không thay thế kiểm tra leak. Kết luận chỉ áp dụng input và ngân sách đã thử; không dùng kết quả hữu hạn để chứng minh mọi đầu vào.',
        'Chuẩn biên dịch C11;': 'Chuẩn biên dịch C11; sanitizer dùng -O1 -g -fsanitize=address,undefined -fno-sanitize-recover=all -fno-omit-frame-pointer. Compiler/gcov dùng toolchain portable, không cần sửa PATH hệ thống. Seed và ngân sách được ghi riêng cho mỗi chiến dịch.',
        'Mỗi ca có mã yêu cầu': 'Excel cập nhật có 5 sheet trọng tâm, 12 yêu cầu và 57 ca thiết kế: 6 ca nghiệp vụ, 43 ca JSON/đường cấp phát, 6 ca guard và 2 ca helper trước/sau vá. Đây không phải 57 ca đều đã chạy Ch4. Kết quả thực thi phải dẫn đúng log và snapshot; không dùng expected thay actual.',
        'Oracle Python được dùng': 'Oracle Python đối chiếu mã trả và trạng thái cuối của C. Guard có hợp đồng mã riêng; characterization chỉ kiểm tra lại hành vi đã ghi nhận. Các ca helper trước/sau dùng harness cố định, không được ghi thành lỗi do chiến dịch fuzz Ch4 tìm thấy.',
        'Ngày 22/10/2026, kiểm tra': 'Ngày 22/10/2026, kiểm tra riêng đường gọi trực tiếp helper đã xác nhận lỗi ghi vượt mảng trong jsonInteger. Đây là kết quả harness bổ trợ, chưa có bằng chứng chiến dịch fuzz Ch4 phát hiện lỗi này. Bản vá và kiểm tra lại được lưu với snapshot trước/sau; số liệu fuzz cũ vẫn thuộc mã lịch sử.',
        'reserve dùng int64_t': 'reserve dùng int64_t cho biểu thức trung gian; guard kiểm tra quantity từ 1 đến 100 trước khi so reserved > INT_MAX - quantity. Các ca biên kiểm tra nhánh từ chối tràn. Đây là cơ chế phòng vệ hiện có, không có cặp lỗi/vá riêng cho CWE-190; chương trình lỗi mẫu chỉ là đối chứng sanitizer.',
        'Sau calloc thành công, jsonEvaluate thu': 'Sau calloc thành công, jsonEvaluate thu mã lỗi, gọi free rồi return. Có test tiêm lỗi cấp phát và cân bằng cấp phát hỗ trợ kiểm tra đường lỗi. Các kiểm tra này không thay thế LeakSanitizer hoặc chứng minh không rò rỉ trên mọi đầu vào.',
        'Người gọi helper vẫn phải': 'Người gọi helper vẫn phải cung cấp buffer hợp lệ có ít nhất size byte và token đúng cú pháp. Kiểm tra độ dài không xác minh được dung lượng thật của con trỏ tùy ý. Chưa kiểm toán toàn bộ helper, chưa có chiến dịch fuzz Ch4 hoặc đo leak sau vá trong hồ sơ; không tự gán CVE.',
        'results/report-final/pytest-combined.log ghi:': 'Có log hồi quy ở mốc mã lịch sử, nhưng số test tổng hợp không được coi là số input fuzz. Phần dưới chỉ trình bày các lượt JSON và sanitizer độc lập, cùng phạm vi snapshot của chúng.',
        'Tất cả lượt trên ghi 0 crash': 'Tất cả lượt trên ghi 0 crash, 0 timeout, 0 vi phạm bất biến và 0 lệch oracle. Coverage-guided có độ phủ cao hơn black-box ở các seed này; kết quả chỉ thuộc ngân sách, seed và snapshot của phép đo.',
        '6.10. Chiến dịch': '6.4. Chiến dịch JSON trong Ch4 của ứng dụng',
        '6.11. Ca kiểm thử': '6.5. Ca kiểm thử tiêu biểu và truy vết kết quả',
        '6.13 Kiểm tra': '6.6. Kiểm tra hồi quy bản vá helper',
        'Excel giữ 55 ca': 'Excel hiện có 57 ca thiết kế và liên kết yêu cầu. Những ảnh dưới là snapshot lịch sử của các ca JSON vẫn được giữ nguyên trong bảng hiện tại; cột KetQuaMongDoi không phải actual hoặc bảng PASS. Sáu ca guard là thiết kế bổ sung, không dùng để tuyên bố đã có một chiến dịch Ch4 tương ứng.',
        'Cố định ngay một mục tiêu/nhánh': 'Cố định đích fuzz, miền input, oracle và cách đo trước mỗi chiến dịch. Tách stock_json, stock_json_case và guard để tránh so coverage của các hàm khác nhau. Lưu seed, snapshot và log của đúng lượt chạy để tái lập.',
        'Đã có so sánh cùng nhánh': 'Kết quả chỉ áp dụng mô hình C và ngân sách đã ghi. Chưa kiểm thử đồng thời/transaction backend hoặc toàn website; không bật LeakSanitizer. Không lấy số liệu lịch sử để xác nhận mã mới sau vá.',
        'Potential possible leak 24 byte': 'Kiểm tra leak chuyên biệt không thuộc nội dung chính của báo cáo Ch4 này; không kết luận hết leak từ ASan/UBSan.',
        'Giữ đúng giới hạn LSan/Dr. Memory': 'Giữ đúng giới hạn sanitizer và phạm vi mô hình C; không cộng các lượt chạy độc lập thành một chiến dịch.',
        'Mô hình kiểm thử đã bao phủ chuỗi': 'Báo cáo tập trung Ch4: từ hợp đồng JSON và đích C đến tạo input, đo coverage, kiểm tra bất biến, oracle và hồi quy. Đặc tả và giới hạn kiểu/byte/record là cơ sở thiết kế dữ liệu fuzz, không được trình bày thành kết quả phân tích của chương khác.',
        'Bằng chứng chính gồm 41 ca': 'Bằng chứng chính gồm các chiến dịch JSON và sanitizer theo từng snapshot, 41 hồi quy JSON và 40 ca DLL. Lượt dài JSON có 3.000 input; lượt sanitizer dài có 1.384 input và không ghi nhận finding. Các bộ này độc lập, không cộng thành một phép chạy mới hoặc coi mọi ca Excel đã PASS.',
        'Lượt GUI JSON ngày 21/10 chạy': 'Lượt Ch4 GUI ngày 21/10 chạy 437 input, giữ chín test và hiển thị coverage hàm 82%. Chín test đều trả 100 do dữ liệu bị từ chối; bằng chứng đường nghiệp vụ hợp lệ dựa thêm trên hồi quy với oracle so sánh trạng thái. Coverage một hàm không đại diện toàn chương trình.',
        'Các chiến dịch lịch sử không ghi nhận': 'Các chiến dịch lịch sử không ghi nhận crash, timeout, vi phạm bất biến hoặc lệch oracle trong ngân sách đã báo cáo. Kiểm tra harness ngày 22/10 xác nhận UBSan finding ở helper gọi trực tiếp với 65 chữ số; cùng ca bị từ chối sau vá. Không suy rằng fuzzing đã phát hiện lỗi đó hoặc mọi đường gọi đều an toàn.',
        'Báo cáo đã có chuỗi bằng chứng': 'Đã có bằng chứng lỗi helper, bản vá và năm nhóm kiểm tra cố định sau vá. Chưa có chiến dịch fuzz Ch4 sau vá trong hồ sơ; coverage lịch sử vẫn thuộc snapshot cũ. Khi đổi quy tắc nghiệp vụ, cần đối chiếu nguồn tham chiếu, đặc tả và ca hồi quy. Mô hình không kiểm thử trực tiếp HTTP, cơ sở dữ liệu hoặc PayOS.',
        'Các summary và log gốc nằm': 'Summary và log gốc được giữ trong từng thư mục results. Ảnh GUI do Phong cung cấp là bằng chứng của lượt được nêu, không thay cho raw input/log. Báo cáo và Excel đã rút phạm vi nội dung về Ch4; workbook không phải mẫu nhập đủ mọi chức năng của ứng dụng.',
        'Runner có thể tạo thư mục': 'Runner tạo thư mục kết quả mới khi chạy. Compiler và .venv không cần đưa vào gói nguồn. Excel là tài liệu thiết kế Ch4 đã biên tập thủ công, không tự sinh lại từ --print-spec để ghi đè vì dữ liệu mẫu trong mã vẫn có phạm vi rộng hơn.',
    }
    for paragraph in doc.paragraphs:
        for prefix, value in changes.items():
            if paragraph.text.startswith(prefix):
                replace_text(paragraph, value)
                break
    for paragraph in doc.paragraphs:
        if 'Mục 6.13 liệt kê' in paragraph.text:
            replace_text(paragraph, paragraph.text.replace('Mục 6.13', 'Mục 6.6'))
        if paragraph.text.startswith('Chuỗi snapshot trước vá'):
            replace_text(paragraph, paragraph.text + ' Đây là hồ sơ harness, không phải kết quả lỗi do chiến dịch fuzz Ch4 phát hiện.')
        if paragraph.text.startswith('Các cơ chế CWE-190, CWE-476 và CWE-401'):
            paragraph._p.getparent().remove(paragraph._p)
    for paragraph in doc.element.xpath('//w:txbxContent//w:p'):
        nodes = paragraph.xpath('.//w:t')
        if ''.join(n.text or '' for n in nodes) == 'BẰNG JSON, SANITIZER VÀ CONCOLIC':
            nodes[0].text = 'BẰNG FUZZING VÀ ĐO ĐỘ PHỦ'
            for node in nodes[1:]:
                node.text = ''
    environment = next(t for t in doc.tables if t.cell(0, 0).text == 'Thành phần')
    for row in list(environment.rows)[1:]:
        name = row.cells[0].text
        if name in ('Z3', 'Dr. Memory'):
            row._tr.getparent().remove(row._tr)
        elif name == 'Hệ điều hành và ABI':
            row.cells[0].text = 'Hệ điều hành'
            row.cells[1].text = 'Windows x64'
            row.cells[2].text = 'Môi trường thực hiện; kết quả gắn với compiler và snapshot của từng lượt.'
        elif name == 'SecLabFramework':
            row.cells[1].text = 'Educational Security Suite; dùng tab Ch4 Coverage & Fuzzing'
            row.cells[2].text = 'Ảnh chiến dịch GUI đã lưu; chưa có định danh bản build chính xác.'
    evidence = next(t for t in doc.tables if t.cell(0, 0).text == 'Thư mục / file')
    for row in list(evidence.rows)[1:]:
        row._tr.getparent().remove(row._tr)
    rows = [
        ('results/stock-json-20261018-083428-95703d48', 'Chiến dịch JSON dài; metadata, snapshot và summary theo seed/chế độ.'),
        ('results/stock-json-20261018-093532-9d53d681', 'Lượt JSON ngắn; 41 hồi quy CLI, 40 DLL, ca nghiệp vụ và trạng thái actual.'),
        ('results/sanitizer-20261018-084125-187616c4', 'Lượt ASan+UBSan dài, 1.384 input; không bật LeakSanitizer.'),
        ('results/sanitizer-20261018-093549-a35b9b72', 'Lượt sanitizer ngắn trên snapshot lịch sử khác.'),
        ('badminton_spec.xlsx', '5 sheet trọng tâm Ch4, 12 yêu cầu, 57 ca thiết kế; không phải 57 ca đã chạy/PASS.'),
        ('results/patch-audit-20261022', 'Harness, snapshot trước/sau, log UBSan và kiểm tra hồi quy bản vá; chưa phải fuzz sau vá.'),
    ]
    for values in rows:
        for cell, text in zip(evidence.add_row().cells, values):
            cell.text = text
    for table in doc.tables[1:]:
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn('w:p'):
            Paragraph(previous, doc).paragraph_format.keep_with_next = True
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1
    figure_map = {}
    for paragraph in doc.paragraphs:
        match = re.match(r'Hình (\d+)\.', paragraph.text)
        if match:
            figure_map[int(match[1])] = len(figure_map) + 1
            replace_text(paragraph, re.sub(r'^Hình \d+\.', 'Hình ' + str(len(figure_map)) + '.', paragraph.text))
    assert doc.tables[0]._tbl.xml == original_group
    doc.core_properties.title = 'Kiểm thử fuzzing mô hình tồn kho Badminton và đo độ phủ'
    doc.save(WORK / 'Baocao.ch4.docx')
    return figure_map


def update_excel_data():
    workbook = load_workbook(WORK / 'badminton_spec.before.xlsx', read_only=True)
    names = ['ThongTin', 'YeuCauChucNang', 'YeuCauPhiChucNang', 'RangBuoc', 'TestCase']
    data = {}
    for name in names:
        data[name] = [list(row) for row in workbook[name].iter_rows(values_only=True) if any(v is not None for v in row)]
    data['ThongTin'] = [
        ['Truong', 'GiaTri'],
        ['name', 'badminton_stock_json_ch4'],
        ['description', 'Đặc tả và thiết kế kiểm thử Ch4 Coverage & Fuzzing cho mô hình tồn kho C; một sản phẩm, tối đa một biến thể.'],
        ['source', 'Nguồn Java đóng băng trong reference_backend; stock_spec.json; generator ghép adapter JSON viết riêng. Không tự dịch toàn bộ Java.'],
        ['input', data['ThongTin'][4][1]],
        ['fields', 'available, reserved, variant_available, variant_reserved: int32. action: 0–15; items chứa record chỉ có quantity.'],
        ['limits', 'Đầu vào ≤16.384 byte; token số ≤64 byte; 1–100 record; quantity mỗi dòng 1–100, tổng ≤100. Giới hạn adapter, không phải HTTP đã xác nhận.'],
        ['target', 'Nạp stock_json_seclab.c vào Ch4. stock_json nhận JSON; stock_json_case tạo fixture từ bốn số; stock_reserve_guard chỉ kiểm tra điều kiện.'],
        ['scope', 'Không HTTP, database, PayOS hoặc đồng thời. Workbook này là tài liệu thiết kế Ch4, không còn mẫu đầy đủ để nhập các chức năng khác của ứng dụng.'],
        ['action', 'Mặt nạ 8 bật biến thể; hai bit thấp: 0 giữ, 1 hoàn, 2/3 tất toán; mặt nạ 4 lặp hoàn/tất toán. Mã fixture, không phải enum HTTP.'],
        ['errors', 'stock_json: 0 thành công; 1–7 nghiệp vụ; 100 JSON/schema/kiểu; 101 độ dài/NULL; 102 record; 103 quantity; 104 fixture; 105 cấp phát; 106 bất biến. Guard dùng mã riêng 0–5.'],
        ['oracle', 'So mã trả và trạng thái kho của C với oracle Python. Không chỉ kiểm tra exit 0 hoặc code 0. Expected trong TestCase không phải actual.'],
        ['execution', '57 ca thiết kế: 6 nghiệp vụ + 43 JSON/đường cấp phát + 6 guard + 2 helper. Không tự đánh dấu cả 57 PASS; xem log của đúng snapshot.'],
        ['fuzz_evidence', 'Lượt JSON lịch sử 083428: 3.000 input, 41 CLI và 40 DLL; lượt 093532: 100 black-box +100 coverage-guided. Xem summary trong results tương ứng.'],
        ['gui_evidence', 'Ảnh Ch4 ngày 21/10: stock_json, 437 input, giữ 9 test đều trả 100, coverage hàm 82%. Không đại diện toàn parser hoặc kết quả sau vá.'],
        ['sanitizer', 'ASan+UBSan lịch sử: lượt dài 1.384 input +41 hồi quy; 0 finding trong tải đã thử. Không đo gcov, không bật LeakSanitizer; không chứng minh hết leak.'],
        ['guard_cases', 'TC-SMT1–6 giữ ID lịch sử, mô tả nhánh 0–5 để thiết kế fuzz guard tại Ch4. Không chuyển kết quả GUI ở tab khác thành chiến dịch Ch4.'],
        ['patch', 'jsonInteger gọi trực tiếp với 65 chữ số: UBSan vượt char[64] trước vá; sau vá trả 0, giữ output=77. Đây là harness cố định; chưa xác nhận fuzz Ch4 đã phát hiện lỗi.'],
        ['patch_source', 'results/patch-audit-20261022/before-results.json; after-results.json; regression-tests.log; verification.json. Chưa có chiến dịch fuzz/GUI sau vá trong hồ sơ.'],
        ['memory', 'calloc 100 phần tử int32_t; kiểm tra NULL và free sau cấp phát thành công. Test cấp phát/cân bằng bổ trợ oracle bộ nhớ, không phải kết luận không leak tuyệt đối.'],
        ['sync', 'Mã runner vẫn có dữ liệu mẫu 10 sheet/55 ca khi dùng --print-spec. Không dùng đầu ra đó ghi đè workbook Ch4 đã biên tập này.'],
    ]
    data['RangBuoc'] = [row for row in data['RangBuoc'] if row[0] != 'CON-6']
    data['RangBuoc'].append(['CON-8', 'Expected, hồi quy cố định, input fuzz và harness bản vá là các loại bằng chứng khác nhau; không cộng số lượng hoặc gán nhầm nguồn.', 'HIGH'])
    data['RangBuoc'].append(['CON-9', 'Sau sửa C phải biên dịch/wrap lại. DLL gốc cũ không chứng minh bản vá; coverage lịch sử chỉ áp dụng snapshot lịch sử.', 'HIGH'])
    for row in data['TestCase'][1:]:
        if str(row[0]).startswith('TC-SMT'):
            row[4] = 'Thiết kế ca cho stock_reserve_guard trong Ch4; nạp C hiện tại, biên dịch/wrap lại; đối chiếu nhánh và mã mong đợi. Không mặc định ca này đã chạy trong chiến dịch Ch4.'
            code = {'TC-SMT1': 1, 'TC-SMT2': 2, 'TC-SMT3': 3, 'TC-SMT4': 4, 'TC-SMT5': 5, 'TC-SMT6': 0}[row[0]]
            row[5] = f'Trả {code}. Guard chỉ kiểm tra điều kiện; mã 0 không cập nhật tồn kho. Đây là expected theo đặc tả, không phải actual Ch4.'
    data['TestCase'] += [
        ['TC-P01', 'REQ-J2', 'Helper từ chối token vượt 64 byte sau vá', 'Gọi trực tiếp jsonInteger qua harness với buffer hợp lệ gồm 65 chữ số 1, size=65, output ban đầu 77; không qua parser JSON.', 'Đọc snapshot/log trước vá và sau vá trong results/patch-audit-20261022; đây là ca cố định, không tự chạy script ghi đè hồ sơ.', 'Trước vá: UBSan index 64 vượt char[64], exit 1. Sau vá: result=0, output=77, exit 0. Không ghi đây là lỗi do chiến dịch fuzz Ch4 phát hiện.', 'HIGH'],
        ['TC-P02', 'REQ-J2, REQ-J5', 'Helper từ chối NULL hoặc độ dài không an toàn', 'Harness sau vá: input NULL, output NULL, size=0 hoặc size=65 với buffer ngắn; output hợp lệ ban đầu 77.', 'Đối chiếu ca bounds trong after-results.json và test_json_integer_rejects_unsafe_direct_inputs. Không dùng con trỏ số tùy ý trong GUI.', 'Từ chối với result=0; giữ output khi output hợp lệ. Chưa chạy các ca bounds trên snapshot trước vá; không dựng actual trước vá.', 'HIGH'],
    ]
    assert len(data['TestCase']) - 1 == 57
    (WORK / 'excel-content.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    return {name: len(rows) - 1 for name, rows in data.items()}


if __name__ == '__main__':
    result = {'figure_map': update_report(), 'sheet_records': update_excel_data(), 'scope': 'Ch4; no fuzz rerun; preface/group table preserved'}
    (WORK / 'scope-audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False))
