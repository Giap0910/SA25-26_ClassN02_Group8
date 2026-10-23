from pathlib import Path
import hashlib
import json
import re
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('C:/Project/Badminton_Stock_Security')
WORK = ROOT/'results/report-final/word-complete-20261021'
BASE = WORK/'Baocao.before.docx'
TARGET = WORK/'Baocao.checked.docx'
DOC = Document(BASE)


def anchor(prefix):
    return next(p for p in DOC.paragraphs if p.text.startswith(prefix))


def add_before(reference, text, style='Normal'):
    p = reference.insert_paragraph_before(text, style)
    if style == 'Normal':
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.12
        for run in p.runs: run.font.size = Pt(14)
    return p


def format_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths): col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:'+side)
        for key, value in [('val','single'), ('sz','4'), ('color','D9D9D9')]: e.set(qn('w:'+key),value)
        borders.append(e)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(table.rows):
        if index == 0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DFE8F0' if index == 0 else 'FFFFFF')
            cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(4)
                for run in p.runs: run.font.size = Pt(11.5); run.bold = index == 0


def add_table(reference, rows, widths):
    table = DOC.add_table(rows=0, cols=len(widths))
    for values in rows:
        for cell, value in zip(table.add_row().cells, values): cell.text = str(value)
    format_table(table, widths)
    reference._p.addprevious(table._tbl)
    add_before(reference, '')
    return table


def add_code(reference, text):
    p = add_before(reference, text)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_together = True
    for run in p.runs: run.font.name = 'Consolas'; run.font.size = Pt(11.5)
    return p


def module_section():
    ref = anchor('Luồng thực hiện')
    add_before(ref, '2.1 Các module của mô hình kiểm thử', 'Heading 2')
    add_before(ref, 'Đối tượng kiểm thử là module tồn kho được mô hình hóa bằng C, gồm giữ hàng, hoàn khi hủy và tất toán. Các thành phần bên dưới hỗ trợ sinh mô hình, đọc dữ liệu và kiểm tra kết quả; chúng không thay thế các dịch vụ đang chạy của website.')
    rows = [['Module và file', 'Đầu vào', 'Đầu ra và trách nhiệm'],
            ['Đặc tả và sinh mô hình\nstock_spec.json\ngenerate_simulator.py', 'Quy tắc, kiểu, miền giá trị và nguồn Java tham chiếu đã đóng băng', 'Sinh nghiệp vụ C và guard từ đặc tả. Không tự dịch toàn bộ Java sang C.'],
            ['Nghiệp vụ tồn kho\nstock_simulator.c\nreserve, settle', 'Trạng thái kho, tổng quantity và action', 'Mã trả và trạng thái kho mới; kiểm tra kho thiếu, tràn số, hủy/tất toán lặp.'],
            ['Đọc JSON và cấp phát\nstock_json_adapter.c\njsonEvaluate', 'Chuỗi byte JSON và độ dài', 'Kiểm tra schema, kiểu, số record; tạo Input; cấp phát và giải phóng mảng quantities.'],
            ['Đích nạp GUI\nstock_json_seclab.c', 'Chuỗi JSON hoặc tham số số nguyên của guard', 'Ghép nghiệp vụ và adapter; stock_json trả mã; stock_reserve_guard chỉ kiểm tra điều kiện.'],
            ['Điều phối fuzz và oracle\nstock_json_fuzz.py\nfuzz_runner.py', 'Fixture, seed, ngân sách và đích C', 'Chạy hồi quy/fuzz, so mã và trạng thái C với oracle Python; lưu log và coverage.'],
            ['Kiểm thử hồi quy\ntest_generator.py\ngói pytest đã xuất', 'Đặc tả, generator, adapter, fixture và test GUI', 'Kiểm tra sinh mã, tính tương đương, đường lỗi bộ nhớ và chạy lại hành vi đã xuất.']]
    add_table(ref, rows, [2.25, 1.8, 2.45])
    add_before(ref, 'Ranh giới cần phân biệt: stock_json chỉ trả mã kết quả cho GUI; CLI của adapter xuất cả trạng thái kho để runner đối chiếu. Guard trả 0 nghĩa là qua kiểm tra điều kiện, chưa thực hiện cập nhật tồn kho.')


def example_section():
    ref = anchor('Những phần đã và chưa thực hiện')
    add_before(ref, '2.5 Ví dụ JSON và luồng xử lý', 'Heading 2')
    add_before(ref, 'Ví dụ dưới lấy từ ca reserve trong hồi quy đã lưu ngày 18/10/2026. Nó mô tả fixture của bộ kiểm thử, không phải request HTTP cho phép người mua đặt giá trị tồn kho.')
    add_code(ref, '{\n  "available": 10, "reserved": 0,\n  "variant_available": 10, "variant_reserved": 0,\n  "action": 0, "items": [{"quantity": 2}]\n}')
    add_table(ref, [['Trường', 'Giá trị và ý nghĩa trong ví dụ'],
                    ['available / reserved', '10 sản phẩm khả dụng, 0 sản phẩm đang giữ chỗ.'],
                    ['variant_available / variant_reserved', '10 và 0 là kho biến thể trong fixture; action=0 không bật nhánh biến thể.'],
                    ['action', '0: chỉ giữ hàng một lần, không hủy hoặc tất toán tiếp.'],
                    ['items / quantity', 'Một record, quantity là số 2; tổng số lượng cần giữ bằng 2.']], [2.25,4.25])
    add_before(ref, 'Reader đọc đúng giới hạn byte, parser xác nhận đủ sáu khóa gốc và kiểu của quantity. Sau khi calloc thành công, parser ghi một quantity vào mảng tối đa 100 phần tử. jsonQuantity kiểm tra số record và tổng, rồi jsonStock gọi reserve. Mảng heap được free trước khi jsonEvaluate trả mã.')
    add_code(ref, '{"code":0,"available":8,"reserved":2,\n "variant_available":10,"variant_reserved":0,"state":1}')
    add_before(ref, 'Đây là actual của ca reserve trong regression.json và khớp oracle: khả dụng giảm 2, giữ chỗ tăng 2, state=1. Không dùng riêng code=0 để thay thế kiểm tra các trường trạng thái.')
    add_before(ref, 'Ví dụ lỗi sai kiểu dùng cùng fixture nhưng đổi quantity từ số 2 thành chuỗi "2". Ca string đã chạy trả code=100; các trường output khởi tạo bằng 0. Parser không ép chuỗi thành số và không chạy nghiệp vụ giữ hàng. Các số 0 này không có nghĩa tồn kho thật trong cơ sở dữ liệu bị xóa.')
    add_code(ref, '"items": [{"quantity": "2"}]\n→ code=100, available=0, reserved=0, state=0')
    add_before(ref, 'Nguồn hai kết quả: results/stock-json-20261018-093532-9d53d681/regression/regression.json, tên ca reserve và string; fixture gốc được lưu cùng thư mục. Đây là diễn giải lại bằng chứng cũ, không phải lượt thực thi mới.')


def environment_section():
    ref = anchor('Truy vết yêu cầu và quy tắc kết luận')
    add_before(ref, '3.3 Môi trường và điều kiện tái lập', 'Heading 2')
    add_before(ref, 'Bảng tập hợp cấu hình có thể xác minh. Phiên bản từ log lịch sử được phân biệt với phiên bản đọc trực tiếp trên máy ngày 21/10/2026; không suy ngược rằng mọi lượt lịch sử đều dùng phiên bản hiện tại.')
    rows = [['Thành phần', 'Cấu hình xác minh', 'Nguồn và phạm vi'],
            ['Hệ điều hành và ABI', 'Windows x64; build 19045 trong log Dr. Memory', 'Log leak ngày 18/10; layout C phụ thuộc ABI này.'],
            ['Python của dự án', '3.13.13, 64-bit AMD64', '.venv/Scripts/python.exe, đọc phiên bản ngày 21/10.'],
            ['pytest', '9.1.1', 'Thư viện trong .venv, đọc ngày 21/10; không gọi số subtest là số hàm test.'],
            ['GCC và gcov', '16.2.0, MinGW-W64 x86_64 UCRT', 'GCC có trong compiler-analysis.json; gcov cùng bộ toolchain được kiểm tra ngày 21/10.'],
            ['Clang và sanitizer', 'Clang 23.1.3, target x86_64-w64-windows-gnu', 'compiler-analysis.json và lệnh hiện tại; address,undefined trong log sanitizer.'],
            ['Z3', '5.3.1', 'ch3-ast-smt-check.json và .venv hiện tại; ràng buộc guard thuần số.'],
            ['Dr. Memory', '2.6.0, Windows x64, leaks_only', 'leak-summary.json; giữ suppression mặc định và potential_errors.'],
            ['SecLabFramework', 'Educational Security Suite, các chức năng Ch1–Ch4 đã dùng', 'Ảnh GUI do người thực hiện cung cấp; chưa có định danh bản build để ghi chính xác.']]
    add_table(ref, rows, [1.4,2.25,2.85])
    add_before(ref, 'Chuẩn biên dịch C11; sanitizer dùng -O1 -g -fsanitize=address,undefined -fno-sanitize-recover=all -fno-omit-frame-pointer. Compiler/gcov dùng bản portable trong toolchain, không cần sửa PATH hệ thống. Seed và ngân sách của từng phép chạy nằm ở các mục 6.1–6.10, không gộp thành một cấu hình chung.')
    add_before(ref, 'Khi chạy lại, giữ đúng snapshot mã, đặc tả, fixture, seed và lệnh trong metadata của từng lượt. Thời gian và độ phủ có thể đổi khi đổi máy hoặc công cụ; chỉ so kết quả trong cùng phép đo. Các lệnh ở phụ lục tạo kết quả mới, không ghi đè log lịch sử để giữ số liệu cũ.')


def conclusion_section():
    ref = anchor('8. Phụ lục')
    add_before(ref, '7.5 Kết luận chung', 'Heading 2')
    add_before(ref, 'Mô hình kiểm thử đã bao phủ chuỗi đặc tả, sinh nghiệp vụ C, đọc JSON, kiểm tra quy tắc tồn kho và đối chiếu trạng thái với oracle. Kiểu và miền giá trị được xác định; giới hạn byte, token và record được tách khỏi size/alignment của struct và vùng heap. Các phép thử bổ sung kiểm tra đường cấp phát thất bại, cân bằng cấp phát và lỗi bộ nhớ bằng công cụ động.')
    add_before(ref, 'Bằng chứng chính gồm 41 ca hồi quy JSON và 40 ca qua DLL; sáu test guard xuất từ GUI chạy lại đạt, kèm 22/22 probe trong lượt concolic. Lượt sanitizer ch123 chạy 41 hồi quy và 100 mutation seed 1 không ghi nhận finding. Dr. Memory kiểm tra 141 fixture lặp 10 vòng, cả 1.410 đầu ra khớp oracle và không báo definite application leak. Các phép chạy này độc lập, không cộng số lượng thành một chiến dịch duy nhất.')
    add_before(ref, 'Lượt GUI JSON ngày 21/10 chạy 437 input, giữ chín test và hiển thị coverage hàm 82%. Chín test đều đi vào nhánh từ chối dữ liệu với mã 100, nên bằng chứng đường nghiệp vụ hợp lệ phải dựa thêm trên hồi quy có trạng thái thực tế. So sánh cùng nhánh return 3 cho thấy SMT cần ít lần gọi C hơn, trong khi lấy mẫu tăng cường biên nhanh hơn về thời gian ở phép thử đã lưu.')
    add_before(ref, 'Kết luận trong phạm vi tải thử là chưa ghi nhận crash, vi phạm bất biến, lệch oracle hoặc definite application leak của mục tiêu ở các lượt được báo cáo. Potential possible leak 24 byte trong ntdll/TLS vẫn còn; không kết luận hệ thống không có leak tuyệt đối. Mô hình không kiểm thử trực tiếp HTTP, cơ sở dữ liệu, giao dịch đồng thời hoặc toàn bộ website.')
    add_before(ref, 'Báo cáo cung cấp bằng chứng kiểm thử cơ chế phòng vệ hiện có, chưa có cặp lỗi mục tiêu và bản vá v0/v1 được xác nhận. Khi quy tắc nghiệp vụ thay đổi, cần đối chiếu nguồn tham chiếu, cập nhật đặc tả và chạy lại các ca liên quan; kết quả hiện tại chỉ áp dụng snapshot đã ghi nhận.')


def normalize_headings():
    mapping = {'Luồng thực hiện':'2.2 Luồng thực hiện', 'Phạm vi đầu vào':'2.3 Phạm vi đầu vào',
               'Hợp đồng dữ liệu và miền giá trị':'2.4 Hợp đồng dữ liệu và miền giá trị',
               'Những phần đã và chưa thực hiện':'2.6 Phạm vi kết luận',
               'Công cụ và ngân sách':'3.2 Công cụ và ngân sách',
               'Truy vết yêu cầu và quy tắc kết luận':'3.4 Truy vết yêu cầu và quy tắc kết luận',
               'Nếu bắt đầu lại, nên thay đổi gì?':'7.1 Điều chỉnh phương pháp khi thực hiện lại',
               'Điều gì chưa đủ bằng chứng?':'7.2 Giới hạn bằng chứng',
               'Bài học phương pháp luận':'7.3 Bài học phương pháp luận',
               'Điều kiện trước khi chốt nộp':'7.4 Điều kiện trước khi chốt nộp'}
    add_before(anchor('Bước 1.'), '3.1 Quy trình kiểm thử', 'Heading 2')
    if 'Caption' not in DOC.styles: DOC.styles.add_style('Caption',WD_STYLE_TYPE.PARAGRAPH)
    for p in DOC.paragraphs:
        text = p.text.strip()
        if text in mapping: p.text = mapping[text]; p.style = 'Heading 2'
        elif text.startswith('CWE-'): p.style = 'Heading 2'
        elif text.startswith('Hình '): p.style = 'Caption'; p.paragraph_format.keep_with_next = True
        elif not text and p.style.name.startswith('Heading'): p.style = 'Normal'
        if text and p.style.name.startswith('Heading'): p.paragraph_format.keep_with_next = True
        if re.match(r'^[1-8]\. ',p.text) and p.style.name == 'Heading 1':
            p.paragraph_format.page_break_before = True
        if p.text.startswith('Vị trí & mô tảstock_'): p.text = p.text.replace('mô tảstock_','mô tả stock_')
        if p.text.startswith('JSON chứa sáu trường gốc bắt buộc và danh sách items'):
            p.text = p.text.replace('JSON chứa sáu trường gốc bắt buộc và danh sách items', 'JSON có sáu khóa gốc bắt buộc, trong đó danh sách items')
    for name in ['Heading 1','Heading 2','Heading 3','Caption']:
        style = DOC.styles[name]; style.font.color.rgb = RGBColor(0,0,0)
        style.paragraph_format.keep_with_next = True
    DOC.styles['Caption'].font.name = 'Times New Roman'
    DOC.styles['Caption'].font.size = Pt(11.5)
    DOC.styles['Caption'].font.bold = True


def update_toc_field():
    instructions = DOC.element.xpath('//w:instrText')
    for field in instructions:
        if field.text and ' TOC ' in field.text: field.text = ' TOC \\o "1-2" \\h \\z '
    settings = DOC.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None: update = OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'),'true')


def table_and_front_layout():
    preface = anchor('1. Lời mở đầu')
    for br in preface._p.xpath('.//w:br[@w:type="page"]'): br.getparent().remove(br)
    front = True
    for p in list(DOC.paragraphs):
        if p._p is preface._p: front = False
        if front and not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:pict|.//*[local-name()="AlternateContent"]'):
            if p._p is not DOC.paragraphs[0]._p: p._p.getparent().remove(p._p)
    for table in DOC.tables:
        for row_index, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            if props.find(qn('w:cantSplit')) is None: props.append(OxmlElement('w:cantSplit'))
            if row_index == 0 and props.find(qn('w:tblHeader')) is None: props.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.style.name.startswith('Heading'): p.style = 'Normal'
    anchor('7.5 Kết luận chung').paragraph_format.page_break_before = True


def repair_command_layout():
    p = anchor('python -B test_generator.py python -B')
    p.text = '\n'.join([r'.\.venv\Scripts\python.exe -B test_generator.py',
                        r'.\.venv\Scripts\python.exe -B stock_json_fuzz.py --iterations 50 --seeds 1 --seconds 30',
                        r'.\.venv\Scripts\python.exe -B stock_json_fuzz.py --sanitizer --iterations 100 --seeds 1',
                        r'.\.venv\Scripts\python.exe -B -m pytest test_generator.py'])
    for run in p.runs: run.font.name = 'Consolas'; run.font.size = Pt(10)
    p.paragraph_format.keep_together = True
    p.paragraph_format.line_spacing = 1.1
    p2 = anchor('.venv/Scripts/python.exe -B -m pytest test_stock_')
    p2.text = r'.\.venv\Scripts\python.exe -B -m pytest test_stock_json_seclab_stock_reserve_guard_fuzz.py'
    for run in p2.runs: run.font.name = 'Consolas'; run.font.size = Pt(10)


def repair_body_flow():
    in_body = False
    in_hashes = False
    for p in DOC.paragraphs:
        if p.text.startswith('1. Lời mở đầu'): in_body = True
        if in_body:
            for br in p._p.xpath('.//w:br[@w:type="page"]'): br.getparent().remove(br)
            if p._p.xpath('.//w:drawing'):
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.keep_together = True
        if p.text.startswith('8.4.'):
            in_hashes = True
            p.paragraph_format.keep_with_next = True
        elif p.text.startswith('8.5.'): in_hashes = False
        elif in_hashes:
            p.text = re.sub(r' (stock_spec.json|generate_simulator.py|stock_json_adapter.c)',r'\n\1',p.text)
            p.paragraph_format.keep_with_next = not p.text.startswith('e5595517')
            for run in p.runs: run.font.name = 'Consolas'; run.font.size = Pt(9.5)
    anchor('Kiểm tra bộ nộp là').paragraph_format.keep_together = True
    format_table(DOC.tables[-1], [3.25, 3.25])


def verify_content():
    before = Document(BASE)
    assert DOC.paragraphs[0]._p.xml == before.paragraphs[0]._p.xml
    assert [c.text for r in DOC.tables[0].rows for c in r.cells] == [c.text for r in before.tables[0].rows for c in r.cells]
    assert len(DOC.inline_shapes) == len(before.inline_shapes)
    text = '\n'.join(p.text for p in DOC.paragraphs)
    assert not re.search(r'\bCh\s*[56]\b|\bChương\s*[56]\b',text,re.I)
    preface = anchor('1. Lời mở đầu')
    nxt = preface._p.getnext()
    assert '2. Tổng quan đề tài' in ''.join(nxt.itertext())
    for value in ['2.1 Các module','2.5 Ví dụ JSON','3.3 Môi trường','7.5 Kết luận chung']:
        assert value in text


if __name__ == '__main__':
    module_section(); example_section(); environment_section(); conclusion_section()
    normalize_headings(); table_and_front_layout(); repair_command_layout()
    repair_body_flow()
    update_toc_field(); verify_content()
    DOC.save(TARGET)
    print(TARGET)
    print('Baseline SHA256:',hashlib.sha256(BASE.read_bytes()).hexdigest())
