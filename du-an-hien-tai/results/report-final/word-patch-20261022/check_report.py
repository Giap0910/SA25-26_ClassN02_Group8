"""Kiểm tra bảo toàn báo cáo và tạo ảnh QA từ PDF Word đã xuất."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import hashlib
import json
from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image, ImageDraw

WORK = Path(__file__).parent
old = Document(WORK / 'Baocao.before.docx')
new = Document(WORK / 'Baocao.updated.docx')
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def textbox(path):
    with ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        return ''.join(xml.xpath('//w:txbxContent//w:t/text()', namespaces=ns))


assert textbox(WORK / 'Baocao.before.docx') == textbox(WORK / 'Baocao.updated.docx'), 'Bìa thay đổi'
assert [[c.text for c in r.cells] for r in old.tables[0].rows] == [[c.text for c in r.cells] for r in new.tables[0].rows]
assert old.paragraphs[2].text == new.paragraphs[2].text == '1. Lời mở đầu'
with ZipFile(WORK / 'Baocao.before.docx') as z1, ZipFile(WORK / 'Baocao.updated.docx') as z2:
    images1 = sorted(hashlib.sha256(z1.read(n)).hexdigest() for n in z1.namelist() if n.startswith('word/media/'))
    images2 = sorted(hashlib.sha256(z2.read(n)).hexdigest() for n in z2.namelist() if n.startswith('word/media/'))
    assert images1 == images2

pdf_path = WORK / 'Baocao.updated.pdf'
reader = PdfReader(pdf_path)
texts = [page.extract_text() or '' for page in reader.pages]
print('PAGES', len(texts))
for index, text in enumerate(texts, 1):
    first = ' '.join(text.split())[:145]
    print(index, first)
print('CHANGED_LOCATIONS')
for index, text in enumerate(texts, 1):
    if any(term in text for term in ['4.1 Lỗi', '5. Bảng', '6.13', '7.5 Kết', '8.4 Phiên', 'SHA-256', 'index 64']):
        print(index, text[:500].replace('\n', ' '))

pdf = pdfium.PdfDocument(str(pdf_path))
for i, page in enumerate(pdf):
    picture = page.render(scale=1.8).to_pil()
    picture.save(WORK / f'page-{i+1:02}.png')
    page.close()
pdf.close()
for start in range(0, len(texts), 8):
    sheet = Image.new('RGB', (1600, 2260), '#DDDDDD')
    draw = ImageDraw.Draw(sheet)
    for j in range(start, min(start+8, len(texts))):
        picture = Image.open(WORK / f'page-{j+1:02}.png').convert('RGB')
        picture.thumbnail((390, 1070))
        x, y = ((j-start) % 4)*400, ((j-start)//4)*1130
        sheet.paste(picture, (x, y+25))
        draw.text((x+5, y+5), str(j+1), fill='black')
    sheet.save(WORK / f'contact-{start//8+1}.png')
print('PRESERVED: cover text, group table, preface heading, all media hashes')
