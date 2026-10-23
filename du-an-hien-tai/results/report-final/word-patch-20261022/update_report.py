"""Cập nhật báo cáo theo hồ sơ bản vá có sẵn; không thực thi kiểm thử."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib
import json
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path('C:/Project/Badminton_Stock_Security')
WORK = Path(__file__).parent
SOURCE = ROOT / 'Baocao.docx'
BACKUP = WORK / 'Baocao.before.docx'
OUTPUT = WORK / 'Baocao.updated.docx'
AUDIT = ROOT / 'results/patch-audit-20261022'
before = json.loads((AUDIT / 'before-results.json').read_text(encoding='utf-8'))
after = json.loads((AUDIT / 'after-results.json').read_text(encoding='utf-8'))
verification = json.loads((AUDIT / 'verification.json').read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


assert before['runs']['overflow']['returncode'] == 1
assert 'index 64 out of bounds' in before['runs']['overflow']['stderr']
assert after['runs']['overflow']['returncode'] == 0
assert after['runs']['overflow']['stdout'].strip() == 'result=0 output=77'
assert verification['tests_run'] == 5 and verification['failures'] == 0
assert verification['errors'] == 0 and verification['skipped'] == 0
assert sha(AUDIT / 'before-stock_json_seclab.c') == before['source_sha256']
assert sha(AUDIT / 'after-stock_json_seclab.c') == after['source_sha256']
for name in ['stock_json_adapter.c', 'stock_json_seclab.c', 'sl_target.c', 'test_generator.py', 'stock_spec.json']:
    assert sha(ROOT / name) == verification['sha256'][name], name
if not BACKUP.exists():
    shutil.copy2(SOURCE, BACKUP)
assert sha(SOURCE) == sha(BACKUP), 'Báo cáo đã đổi sau khi tạo bản sao; dừng để bảo vệ chỉnh sửa.'
doc = Document(BACKUP)


def anchor(prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))


def para(ref, text, style='Normal'):
    p = ref.insert_paragraph_before(text, style)
    p.paragraph_format.space_after = Pt(7)
    if style == 'Normal':
        p.paragraph_format.line_spacing = 1.12
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
    else:
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def code(ref, text):
    p = para(ref, text)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_together = True
    for r in p.runs:
        r.font.name = 'Consolas'
        r.font.size = Pt(10.5)
    return p


def table(ref, rows, widths):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + side)
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + k), v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.text = str(value)
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DFE8F0' if i == 0 else ('F5F7F9' if i % 2 == 0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(11)
                    r.bold = i == 0
    ref._p.addprevious(t._tbl)
    para(ref, '')
    return t


def remove_between(start, end):
    node = start._p
    while node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


ref = anchor('6. Bằng chứng kiểm chứng')
remove_between(anchor('4. Phân tích lỗ hổng'), ref)
para(ref, '4. Phân tích lỗ hổng và bản vá', 'Heading 1')
para(ref, 'Ngày 22/10/2026, kiểm tra đường gọi trực tiếp helper đã xác nhận lỗi ghi vượt mảng trong jsonInteger. Bản vá và phép kiểm tra lại được lưu riêng với snapshot trước/sau. Các kết quả fuzz, CTL và leak ngày 18–21/10 vẫn thuộc phiên bản lịch sử, không tự chuyển thành kết quả của mã sau vá.')
para(ref, '4.1 Lỗi ghi vượt mảng trong jsonInteger', 'Heading 2')
para(ref, 'Vị trí và điều kiện tái hiện', 'Heading 3')
para(ref, 'Hàm jsonInteger trong stock_json_adapter.c được ghép vào stock_json_seclab.c và sl_target.c. Hàm dùng mảng cục bộ digits gồm 64 phần tử. Snapshot trước vá nằm trong before-stock_json_seclab.c; vòng ghi ở dòng 147 truy cập digits[count++] mà chưa tự kiểm tra độ dài đầu vào.')
para(ref, 'Đường gọi JSON thông thường đi qua jsonScanNumber, vốn đã giới hạn token số tối đa 64 byte. Tuy nhiên shim có export sl_wrap_jsonInteger để gọi trực tiếp helper. Đường gọi này bỏ qua kiểm tra độ dài của parser. Ca tái hiện dùng buffer hợp lệ chứa 65 chữ số 1, size=65 và biến output ban đầu bằng 77. Harness include nguyên snapshot trước vá rồi gọi trực tiếp cùng helper.')
para(ref, 'Nguyên nhân gốc rễ và phạm vi ảnh hưởng', 'Heading 3')
para(ref, 'Hàm phụ thuộc vào điều kiện đã được caller kiểm tra, trong khi công cụ wrapper cho phép gọi helper độc lập. Ở lần ghi thứ 65, count=64 nên truy cập vượt digits[0..63]. Phân loại là CWE-787, ghi ngoài biên bộ nhớ ngăn xếp. Bằng chứng xác nhận lỗi ở helper gọi trực tiếp; chưa xác nhận khả năng gây lỗi qua stock_json, HTTP hoặc backend Java, và chưa đánh giá khả năng khai thác thực thi mã.')
para(ref, 'Bằng chứng trước vá', 'Heading 3')
code(ref, '/* Trích vòng lặp trước vá */\nchar digits[64];\nwhile (pos < size && text[pos] != \'e\' && text[pos] != \'E\') {\n    if (text[pos] == \'.\') afterPoint = 1;\n    else { digits[count++] = text[pos]; fraction += afterPoint; }\n    ++pos;\n}')
para(ref, 'Clang 23.1.3 biên dịch snapshot với ASan+UBSan, -O1 -g và -fno-sanitize-recover=all. Ca overflow thoát với mã 1. Trích chẩn đoán gốc trong before-overflow.log:')
code(ref, "before-stock_json_seclab.c:147:16:\nruntime error: index 64 out of bounds for type 'char[64]'\nSUMMARY: UndefinedBehaviorSanitizer: undefined-behavior")
para(ref, 'UBSan phát hiện vượt chỉ số và dừng tiến trình trước; không diễn giải log này thành thông báo stack-buffer-overflow do ASan. Các ca valid và public trong hồ sơ trước vá đều thoát 0. Ca public xác nhận JSON có token 65 chữ số đã bị từ chối với mã 100 trước khi vào helper.')
para(ref, 'Nội dung bản vá', 'Heading 3')
para(ref, 'Adapter chính được bổ sung điều kiện ngay đầu hàm, trước truy cập text[0]. Mã đích stock_json_seclab.c và sl_target.c được sinh lại từ adapter đã sửa. Kiểm tra mới từ chối con trỏ NULL, đầu vào rỗng hoặc độ dài vượt 64 byte, đồng thời giữ nguyên biến output khi từ chối theo các ca đã kiểm tra.')
code(ref, 'static int jsonInteger(const char *text, size_t size,\n                       int32_t *output) {\n    if (!text || !output || size == 0 || size > 64) return 0;\n    /* Phần chuyển đổi số tiếp tục như trước. */\n}')
para(ref, 'test_generator.py bổ sung kiểm thử helper và kiểm tra tính đồng bộ giữa bộ sinh với mã nộp; phép đối chiếu cho phép chú thích công dụng đầu file nhưng vẫn kiểm tra thân mã. Đặc tả, nguồn Java tham chiếu và quy tắc giữ/hoàn/tất toán không đổi trong bản vá này.')
para(ref, 'Bằng chứng sau vá', 'Heading 3')
para(ref, 'Cùng ca overflow trên snapshot sau vá trả result=0, output=77, exit=0 và stderr rỗng trong after-results.json. Giá trị 0 ở đây nghĩa là helper từ chối chuyển đổi; không phải mã thành công của nghiệp vụ giữ hàng. Các ca valid, public và bounds đều thoát 0.')
code(ref, '/* Trích stdout sau vá của cùng ca 65 chữ số */\nresult=0 output=77\n/* Exit code: 0; stderr rỗng. */')
para(ref, 'Năm phương thức unittest chọn lọc đạt trong 4,140 giây, không failures/errors/skipped. Tập này gồm kiểm tra biên helper, 41 ca JSON CLI và 40 ca DLL, đồng bộ bộ sinh, tiêm lỗi cấp phát và cân bằng cấp phát. DLL shim được biên dịch riêng sau vá xác nhận sáu ca guard đúng mã 0–5 và helper từ chối chuỗi 65 byte. Đây là kiểm tra cố định, chưa phải chiến dịch fuzz sau vá.')
para(ref, 'Nguồn đối chiếu: results/patch-audit-20261022/before-results.json, after-results.json, before-overflow.log, regression-tests.log và verification.json. Mục 6.13 liệt kê kết quả, mục 8.4 lưu hash. Tham chiếu phân loại CWE-787 được giữ trong danh mục tài liệu tham khảo.')
para(ref, '4.2 Các cơ chế phòng vệ khác đã kiểm tra', 'Heading 2')
para(ref, 'Tràn số nguyên CWE-190', 'Heading 3')
para(ref, 'reserve dùng int64_t cho biểu thức trung gian; guard kiểm tra quantity từ 1 đến 100 trước khi so reserved > INT_MAX - quantity. Sáu ca guard bao gồm nhánh từ chối tràn tổng và biến thể. Đây là cơ chế hiện có được kiểm tra, chưa có cặp lỗi/vá riêng cho CWE-190. Mã mẫu UBSan overflow chỉ là đối chứng công cụ.')
para(ref, 'Giải tham chiếu NULL CWE-476', 'Heading 3')
para(ref, 'jsonEvaluate kiểm tra đầu vào và kết quả calloc; cấp phát thất bại trả 105 trước khi dùng mảng quantities. Test tiêm thất bại cấp phát đạt. Bản vá helper bổ sung kiểm tra NULL riêng cho text và output, nhưng các ca NULL không được chạy trên snapshot trước vá; vì vậy không ghi thành một lỗi NULL đã có bằng chứng before/after độc lập.')
para(ref, 'Vòng đời bộ nhớ CWE-401', 'Heading 3')
para(ref, 'Sau calloc thành công, jsonEvaluate thu mã lỗi, gọi free rồi return. Test cân bằng cấp phát đạt. Lượt Dr. Memory lịch sử kiểm tra 141 fixture lặp 10 vòng, tổng 1.410 lần gọi, không báo definite application leak; còn potential possible leak 24 byte ở ntdll/TLS, cùng dấu vết ở đối chứng sạch. Kết quả này thuộc snapshot lịch sử và không thay thế kiểm tra leak sau vá.')
para(ref, '4.3 Giới hạn xác minh bản vá', 'Heading 2')
para(ref, 'Người gọi helper vẫn phải cung cấp buffer hợp lệ có ít nhất size byte và token đúng cú pháp theo hợp đồng. Kiểm tra độ dài không xác minh được dung lượng thật của con trỏ tùy ý. Bản vá chưa kiểm toán toàn bộ helper khác được xuất bởi công cụ; chưa chạy fuzz, GUI, LeakSanitizer hoặc BMC sau vá. Không tự gán CVE.')
para(ref, 'DLL mới patched-seclab-wrap.dll nằm trong thư mục bằng chứng. DLL gốc libstock_json_seclab_wrap.dll chưa được ghi đè nên wrapper Python mặc định có thể vẫn dùng mã cũ. Khi trình diễn trên SecLabFramework, cần nạp lại C đã vá và biên dịch/wrap lại; không dùng DLL cũ để chứng minh bản vá.')

para(ref, '5. Bảng tổng hợp trước và sau vá', 'Heading 1')
para(ref, 'Bảng dưới đối chiếu các ca cố định của bản vá ngày 22/10/2026. Nhãn “chưa chạy trước vá” được giữ khi không có phép đo tương ứng; kết quả lịch sử không được dùng để giả lập một lượt trước vá mới.')
table(ref, [
    ['Ca kiểm tra', 'Trước vá', 'Sau vá'],
    ['Helper nhận 65 chữ số, size=65', 'UBSan báo index 64 vượt char[64]; exit 1.', 'Trả 0, output=77; exit 0; stderr rỗng.'],
    ['Số hợp lệ, int32 biên, số thập phân/số mũ nguyên trong harness', 'Ca valid đạt, exit 0.', 'Ca valid đạt, exit 0; đúng giá trị mong đợi.'],
    ['Token 64 byte biểu diễn số 1 trong harness', 'Đạt trong ca valid.', 'Đạt trong ca valid; kiểm tra helper bổ sung có biên 63/64 byte.'],
    ['NULL input/output, size=0, size=65 với buffer ngắn', 'Chưa chạy các ca bounds trước vá.', 'Ca bounds từ chối an toàn; exit 0.'],
    ['JSON công khai chứa token 65 chữ số', 'stock_json trả 100; ca public đạt.', 'Giữ mã 100; ca public đạt.'],
    ['41 ca JSON CLI và 40 ca DLL', 'Có log lịch sử riêng; không chạy mới trước vá.', 'Đạt trong bộ kiểm tra chọn lọc sau vá.'],
    ['Sáu ca guard qua DLL shim mới', 'Không chạy mới trước vá.', '6/6 đúng mã 0–5.'],
    ['Helper qua DLL shim mới', 'Harness trước vá gọi cùng helper trực tiếp.', 'Từ chối chuỗi 65 byte.']
], [2.25, 2.1, 2.15])
para(ref, 'Kết luận đối chiếu', 'Heading 2')
para(ref, 'Chuỗi snapshot trước vá → chẩn đoán UBSan → kiểm tra đầu hàm → cùng ca được từ chối sau vá đã có bằng chứng cho CWE-787 ở đường gọi trực tiếp helper. Các kiểm tra cố định bổ sung xác nhận hành vi được kỳ vọng trong phạm vi tải thử. Kết quả này không chứng minh toàn bộ chương trình an toàn hoặc website từng bị khai thác.')
para(ref, 'Các cơ chế CWE-190, CWE-476 và CWE-401 được trình bày ở mục 4.2 là phòng vệ đã kiểm tra. Chúng không được gộp thành ba lỗ hổng khác đã vá. Lượt sanitizer tái hiện lỗi helper dùng ASan+UBSan, nhưng chẩn đoán quan sát được trước vá là UBSan.')

# Đồng bộ phát biểu cũ có liên quan trực tiếp đến tình trạng bản vá.
anchor('Tài liệu đề tài cho phép').text = 'Báo cáo trình bày kết quả kiểm thử có ngân sách và bản vá helper xác nhận ngày 22/10/2026. Bằng chứng trước/sau áp dụng jsonInteger khi gọi trực tiếp; các số liệu lịch sử được giữ theo snapshot riêng. Không quy kết lỗi helper thành lỗi HTTP hoặc toàn bộ website.'
anchor('Nếu yêu cầu before/after:').text = 'Cặp snapshot và log trước/sau của lỗi helper đã được lưu tại results/patch-audit-20261022. Trước khi trình diễn bản vá trong GUI, cần nạp lại C và biên dịch/wrap lại để tránh dùng DLL gốc chưa cập nhật.'
anchor('Kết luận trong phạm vi tải thử').text = 'Các chiến dịch lịch sử không ghi nhận crash, vi phạm bất biến, lệch oracle hoặc definite application leak của mục tiêu trong ngân sách đã báo cáo. Kiểm tra riêng ngày 22/10/2026 xác nhận UBSan finding ở helper jsonInteger gọi trực tiếp với 65 chữ số; cùng ca bị từ chối sau vá. Vì vậy không mở rộng kết luận “không tìm thấy lỗi” của các lượt fuzz cũ sang mọi đường gọi. Potential possible leak 24 byte trong ntdll/TLS vẫn chưa được kết luận triệt để.'
anchor('Báo cáo cung cấp bằng chứng kiểm thử').text = 'Báo cáo đã có chuỗi bằng chứng lỗi mục tiêu và bản vá cho đường gọi helper, cùng năm nhóm kiểm tra chọn lọc sau vá đạt. Chưa chạy chiến dịch fuzz, GUI hoặc leak sau vá; số liệu coverage và Dr. Memory lịch sử vẫn thuộc snapshot cũ. Mô hình không kiểm thử trực tiếp HTTP, cơ sở dữ liệu, giao dịch đồng thời hoặc toàn bộ website. Khi quy tắc nghiệp vụ đổi, cần đối chiếu nguồn tham chiếu, cập nhật đặc tả và kiểm tra lại các ca liên quan.'

ref = anchor('7. Phản biện')
para(ref, '6.13 Kiểm tra bản vá helper ngày 22 tháng 10 năm 2026', 'Heading 2')
para(ref, 'Hồ sơ tại results/patch-audit-20261022 lưu snapshot, harness, lệnh biên dịch, phiên bản compiler, exit code và stdout/stderr. Các bảng dưới tổng hợp log đã lưu; việc cập nhật báo cáo không thực thi lại kiểm thử hoặc tạo số đo mới.')
table(ref, [
    ['Nhóm bằng chứng', 'Kết quả đã lưu', 'Giới hạn'],
    ['Harness trước vá', 'overflow: exit 1, UBSan vượt char[64]; valid/public: exit 0.', 'Gọi trực tiếp helper; không có chẩn đoán ASan riêng.'],
    ['Harness sau vá', 'overflow/valid/public/bounds: exit 0; overflow giữ output=77.', 'Ca cố định dưới ASan+UBSan.'],
    ['Bộ unittest chọn lọc', '5 test, 0 failures/errors/skipped; 4,140 giây.', 'Chưa chạy toàn bộ suite hoặc fuzz sau vá.'],
    ['Hồi quy trong bộ chọn lọc', '41 ca JSON CLI và 40 ca DLL đạt.', 'Nằm trong nhóm test, không cộng thành số phương thức unittest.'],
    ['DLL shim sau vá', '6 guard đúng mã; helper từ chối chuỗi 65 byte.', 'DLL riêng trong hồ sơ; DLL gốc chưa đổi.']
], [1.6, 2.6, 2.3])
para(ref, 'Năm test chọn lọc gồm: test_json_integer_rejects_unsafe_direct_inputs, test_json_adapter_and_wrap, test_submission_c_matches_generator, test_json_allocation_failure và test_json_allocation_balance. Nguồn là regression-tests.log và verification.json; đầu ra harness nằm trong before-results.json và after-results.json. Không cộng số test hoặc thời gian này vào chiến dịch fuzz ngày 18/10.')

ref = anchor('8.3. Tài liệu tham khảo')
para(ref, 'Đối với bản vá ngày 22/10, hãy nạp lại stock_json_seclab.c và yêu cầu ứng dụng biên dịch/wrap lại. libstock_json_seclab_wrap.dll ở thư mục gốc chưa được thay thế; DLL đã kiểm tra được lưu riêng là results/patch-audit-20261022/patched-seclab-wrap.dll. Chưa có bằng chứng GUI sau vá. Chỉ gọi helper khi công cụ hỗ trợ truyền buffer, size và con trỏ output đúng hợp đồng; nếu GUI không hỗ trợ, dùng hồ sơ harness và ghi rõ phạm vi.')

# Giữ hash lịch sử trong mục cũ, bổ sung nhãn phiên bản và bảng hash sau vá.
ref = anchor('8.5. Danh mục đối chiếu bằng chứng')
start = anchor('8.4. Danh mục bằng chứng và hash mã nộp')
old_hashes = {}
paragraphs = doc.paragraphs
for i, p in enumerate(paragraphs):
    if p is start:
        break
old_hashes = {
    'stock_json_seclab.c': 'dae8dd9162d9e8929f9a67af7cee1555b90ff42ae70995802c0b89ea185a82c5',
    'stock_spec.json': 'e078aaccd1b2aa5b2b425c09e4292a6c0cc9488c41d8e042451217efe36bb88c',
    'generate_simulator.py': 'c5f31423ba6fa0f8e5f117dacbbc1a3cbb8b7c41338b9d74d3ecd03add097bbd',
    'stock_json_adapter.c': 'e559551769049484cb82874ea724b270a260b84a59d000881174e8d1f2b5a915'
}
remove_between(start, ref)
para(ref, '8.4 Phiên bản và hash mã nguồn', 'Heading 2')
para(ref, 'SHA-256 nhận diện toàn bộ byte của từng file, kể cả chú thích và xuống dòng. Bảng lịch sử được giữ từ bản báo cáo trước cập nhật, không dùng thay cho hash file sau vá. Mỗi hash 64 ký tự được xuống dòng ở giữa để trình bày; nối hai dòng liên tiếp sẽ được giá trị đầy đủ.')


def hash_text(value):
    return value[:32] + '\n' + value[32:]


para(ref, 'Hash đã ghi trong báo cáo trước cập nhật', 'Heading 3')
table(ref, [['File', 'SHA-256 lịch sử']] + [[name, hash_text(value)] for name, value in old_hashes.items()], [2.3, 4.2])
para(ref, 'Các kết quả ngày 18–21/10 phải đối chiếu với mã snapshot trong từng thư mục results và metadata tương ứng. Không mặc định mọi lượt lịch sử dùng cùng hash ở bảng này. Hash generator hiện tại có thể khác do chú thích; khác hash không tự chứng minh có thay đổi nghiệp vụ.')
para(ref, 'Snapshot tái hiện lỗi và snapshot sau vá', 'Heading 3')
table(ref, [
    ['Snapshot trong patch-audit-20261022', 'SHA-256 đã kiểm tra'],
    ['before-stock_json_seclab.c', hash_text(before['source_sha256'])],
    ['after-stock_json_seclab.c', hash_text(after['source_sha256'])]
], [2.3, 4.2])
para(ref, 'Snapshot trước vá đã có chú thích đầu file nên hash có thể khác mốc C lịch sử; nó là nguồn đúng của before-probe.c và log UBSan tái hiện. Snapshot sau vá có hash trùng stock_json_seclab.c hiện tại. Lệnh compiler và kết quả được gắn với từng snapshot trong before-results.json và after-results.json.')
para(ref, 'Hash file hiện tại sau vá ngày 22 tháng 10 năm 2026', 'Heading 3')
names = ['stock_json_seclab.c', 'stock_json_adapter.c', 'sl_target.c', 'generate_simulator.py', 'test_generator.py', 'stock_spec.json', 'results/patch-audit-20261022/patched-seclab-wrap.dll']
table(ref, [['File sau vá hoặc đầu vào đối chiếu', 'SHA-256 hiện tại']] + [[name.replace('results/patch-audit-20261022/', 'patch-audit-20261022/'), hash_text(sha(ROOT / name))] for name in names], [2.3, 4.2])
para(ref, 'Các hash mã đã vá khớp verification.json. Hash DLL được đọc từ file đã lưu; không có bước biên dịch lại trong lần biên soạn này. DLL gốc chưa cập nhật không nằm trong danh sách mã thực thi đã vá. Các hash dài nhận diện byte, không thay thế kiểm thử hành vi.')

# Dùng cỡ chữ monospace cho ô hash để giữ đủ 32 ký tự trên mỗi dòng.
for t in doc.tables:
    if t.cell(0, 1).text.startswith('SHA-256'):
        for row in t.rows[1:]:
            for p in row.cells[1].paragraphs:
                for r in p.runs:
                    r.font.name = 'Consolas'
                    r.font.size = Pt(10)

# Cập nhật danh mục bằng chứng mà không đụng các bảng người dùng tự điền.
evidence_table = next(t for t in doc.tables if t.cell(0, 0).text == 'Thư mục / file')
new_row = evidence_table.add_row()
for cell, value in zip(new_row.cells, ['results/patch-audit-20261022', 'Snapshot trước/sau, harness, log UBSan, năm nhóm kiểm tra chọn lọc, verification.json và DLL shim sau vá.']):
    cell.text = value
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11.5)

doc.save(OUTPUT)

# Kiểm tra bảo toàn nội dung được người dùng yêu cầu giữ nguyên.
original = Document(BACKUP)
updated = Document(OUTPUT)
assert original.tables[0]._tbl.xml == updated.tables[0]._tbl.xml
assert original.paragraphs[0]._p.xml == updated.paragraphs[0]._p.xml
assert anchor('1. Lời mở đầu').text == '1. Lời mở đầu'
with ZipFile(BACKUP) as z1, ZipFile(OUTPUT) as z2:
    media = [n for n in z1.namelist() if n.startswith('word/media/')]
    assert all(z1.read(n) == z2.read(n) for n in media)
    assert z1.read('word/header1.xml') == z2.read('word/header1.xml') if 'word/header1.xml' in z1.namelist() else True
body = '\n'.join(p.text for p in updated.paragraphs)
assert 'chưa có cặp lỗi mục tiêu và bản vá' not in body
assert not any(x in body.lower() for x in ['chương 5', 'chương 6', 'ch5', 'ch6'])
print(json.dumps({'staging': str(OUTPUT), 'backup_sha256': sha(BACKUP), 'staging_sha256': sha(OUTPUT), 'images_preserved': len(media), 'paragraphs': len(updated.paragraphs), 'tables': len(updated.tables)}, ensure_ascii=False, indent=2))
